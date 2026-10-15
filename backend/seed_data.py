import os
import sys
from pathlib import Path
import openpyxl

if sys.stdout:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from backend.database import get_db_connection, init_db

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = PROJECT_ROOT / "data" / "events_seed.xlsx"

def seed_events():
    """Seeds the events table from data/events_seed.xlsx."""
    # Ensure tables exist
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if table already has rows
    cursor.execute("SELECT COUNT(*) AS count FROM events;")
    row_count = cursor.fetchone()["count"]
    if row_count > 0:
        print(f"Skipping seed: events table already contains {row_count} events.")
        conn.close()
        return

    # Check if excel file exists
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(
            f"Seed file not found at '{EXCEL_PATH}'. "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    print(f"Loading seed data from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(str(EXCEL_PATH))
    if "Events" not in wb.sheetnames:
        raise ValueError(f"Sheet 'Events' not found in {EXCEL_PATH}. Available sheets: {wb.sheetnames}")

    sheet = wb["Events"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Spreadsheet 'Events' is empty.")

    header = rows[0]
    expected_header = ['id', 'title', 'category', 'description', 'event_date', 'event_time', 'venue', 'capacity', 'fee', 'icon']
    # Check header columns (normalized to string and lower)
    header_clean = [str(col).strip().lower() if col is not None else "" for col in header]
    if header_clean[:len(expected_header)] != expected_header:
        print(f"Warning: Excel header {header_clean} differs slightly from expected {expected_header}, mapping by position.")

    inserted_count = 0
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = row[:10]

        # Convert date to string format YYYY-MM-DD if openpyxl parsed as datetime
        if hasattr(event_date, "strftime"):
            event_date = event_date.strftime("%Y-%m-%d")
        elif event_date is not None:
            event_date = str(event_date).strip()

        # Convert time to string
        if hasattr(event_time, "strftime"):
            event_time = event_time.strftime("%I:%M %p")
        elif event_time is not None:
            event_time = str(event_time).strip()

        cursor.execute("""
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            int(event_id),
            str(title).strip(),
            str(category).strip(),
            str(description).strip() if description else "",
            event_date,
            event_time,
            str(venue).strip() if venue else "",
            int(capacity),
            int(fee) if fee is not None else 0,
            str(icon).strip() if icon else ""
        ))
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {EXCEL_PATH}.")

def display_events():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events ORDER BY id ASC;")
    rows = cursor.fetchall()
    print("\n--- Current Events in Database ---")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} | {r['category']} | {r['event_date']} {r['event_time']} | {r['venue']} | Cap: {r['capacity']} | Fee: ₹{r['fee']}")
    conn.close()

if __name__ == "__main__":
    seed_events()
    display_events()
